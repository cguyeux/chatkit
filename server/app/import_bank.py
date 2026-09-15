# app/import_bank.py
"""Offline: turn the trainers' reviewed spreadsheet (from generate_bank.py)
into the validated bank. Only rows whose « statut » is validé / ok / oui are
kept; trainers may have edited the question, the choices or the answer, so
each row is re-normalised and re-identified.

  python import_bank.py bank_reviewed.xlsx                 -> question_bank.json (bundled at next build)
  python import_bank.py bank_reviewed.xlsx --push          -> also bank/questions.json in the state store (live, no rebuild)"""
from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.bank import BUNDLED_BANK, STORE_KEY, VALID_STATUSES  # noqa: E402
from app.schemas import normalize_question  # noqa: E402


def read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() == ".json":
        return json.loads(path.read_text(encoding="utf-8"))
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h or "").strip() for h in rows[0]]
    out = []
    for r in rows[1:]:
        d = {headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))}
        out.append(d)
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("sheet")
    ap.add_argument("--out", default=str(BUNDLED_BANK))
    ap.add_argument("--push", action="store_true", help="also write to the state store (needs STATE_S3_* or STATE_LOCAL_DIR)")
    args = ap.parse_args()
    rows = read_rows(Path(args.sheet))
    bank: list[dict] = []
    for r in rows:
        status = str(r.get("statut") or r.get("status") or "").strip().lower()
        if status not in VALID_STATUSES:
            continue
        kc_id = str(r.get("kc_id") or "").strip()
        raw = {
            "text": r.get("question") or r.get("text"),
            "choices": [r.get("A"), r.get("B"), r.get("C"), r.get("D")] if "A" in r else r.get("choices"),
            "answer": r.get("reponse") or r.get("answer"),
            "explanation": r.get("justification") or r.get("explanation") or "",
            "page": r.get("page") or None,
        }
        clean = normalize_question(raw, kc_id, shuffle=False)
        if not clean:
            print(f"skipped (invalid): {str(raw.get('text'))[:80]}", file=sys.stderr)
            continue
        clean["difficulty"] = str(r.get("difficulte") or r.get("difficulty") or "")
        clean["status"] = "validé"
        clean["reviewer_comment"] = str(r.get("commentaire") or "")
        bank.append(clean)
    Path(args.out).write_text(json.dumps(bank, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(bank)} validated questions -> {args.out}")
    if args.push:
        from app.storage import store

        store().put_json(STORE_KEY, bank)
        print(f"pushed to state store key {STORE_KEY}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
